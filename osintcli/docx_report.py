"""Editable DOCX from the same snapshot as PDF/JSON; no remote media."""
import json
from .core import AppError, atomic_path, clean


def docx_export(data,path):
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:raise AppError("DEPENDENCY","DOCX membutuhkan python-docx; jalankan installer profil reports.")
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.15
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].font.name='DejaVu Sans'
        doc.styles[name].font.color.rgb=RGBColor.from_string('006B78')
    for style in doc.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
    doc.core_properties.author='R_HMT';doc.core_properties.title='OSINTCLI / '+data['case']['id']
    def text(value):
        # XML 1.0 cannot contain noncharacters or control bytes.
        return ''.join(c for c in clean(value) if c in '\n\t' or (32<=ord(c)<=0xD7FF) or (0xE000<=ord(c)<=0xFFFD) or (0x10000<=ord(c)<=0x10FFFF))
    def para(value,style=None):return doc.add_paragraph(text(value),style)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for cell,value in zip(t.rows[0].cells,headers):cell.text=text(value)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=text(value)
        return t
    case=data['case']
    doc.add_heading('OSINTCLI',0);para('PUBLIC-SOURCE INTELLIGENCE / by R_HMT')
    doc.add_heading(case['target'],1)
    para('Case '+case['id']+' / '+case['state'])
    para('Dibuat: '+data['generated_at']+'\nGenerator: '+data['generator'])
    if 'demo' in case['scope']:para('OFFLINE FIXTURE / DEMONSTRATION')
    doc.add_heading('Ringkasan',1)
    para(f"{data['counts']['total']} bukti terkumpul. {len(data['relationships'])} hubungan publik tercatat. Kesamaan username tidak membuktikan identitas atau kepemilikan akun.")
    table(['Status','Jumlah'],[[k,str(v)] for k,v in sorted(data['counts']['by_status'].items())])
    doc.add_heading('Metode dan cakupan',1)
    para('Modul: '+', '.join(case['scope']))
    para(data['methodology'])
    table(['Modul','Status'],[[j['adapter'],j['state']] for j in data['jobs']])
    doc.add_heading('Hubungan dan review',1)
    if not data['relationships']:para('Belum ada tautan lintas profil yang diamati.')
    for e in data['relationships']:
        para(e['from']+'\n→ '+e['to'])
        para(e['relation']+' / depth '+str(e['depth'])+' / '+e['review']['decision'])
        para('Sumber: '+e['source_url']+'\nWaktu: '+e['collected_at'])
        if e['review'].get('note'):para('Catatan: '+e['review']['note'])
    doc.add_heading('Register bukti',1)
    rows=data['evidence'][:15] if data['summary_only'] else data['evidence']
    if data['summary_only']:para('Ringkasan maksimal 15 bukti; JSON/CSV menyimpan seluruh register.')
    reviews={r['evidence_id']:r for r in data['reviews']}
    for i,r in enumerate(rows,1):
        doc.add_heading(f"{i:03d} / {r['title']}",2)
        para(r['status']+' / '+r['id']+' / '+r['collected_at'])
        try:value=json.loads(r['normalized_value'])
        except (ValueError,TypeError):value=r['normalized_value']
        if isinstance(value,dict):
            for k,v in value.items():para(str(k)+': '+str(v))
        else:para(r['normalized_value'])
        para('Sumber: '+(r['source_url'] or r['source_name']))
        para('Alasan: '+r['verification_reason'])
        if r['id'] in reviews:para('Review: '+reviews[r['id']]['decision']+' / '+reviews[r['id']]['note'])
        para('SHA-256: '+r['digest'])
    doc.add_heading('Keterbatasan dan jejak',1)
    para('FAILED_SOURCE berarti sumber gagal dibaca, bukan akun tidak ada. Metadata dan tautan dapat berubah. Review accepted adalah penilaian operator, bukan pembuktian identitas oleh sistem. Nomor kontak yang dipublikasikan tidak memastikan nomor aktif atau terdaftar WhatsApp.')
    para('Dataset SHA-256: '+data['dataset_digest'])
    if data['narrative']:
        doc.add_heading('Narasi AI yang ditinjau',1);para(data['narrative']['text'])
    footer=sec.footer.paragraphs[0]
    footer.add_run('OSINTCLI / '+case['id']+' / ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    with atomic_path(path) as tmp:doc.save(tmp)
